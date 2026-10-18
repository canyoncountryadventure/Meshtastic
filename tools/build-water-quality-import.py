#!/usr/bin/env python3
"""Convert the Southeast Utah workbook into MLID-keyed JSON records for import."""
import argparse,json
from datetime import datetime,timezone,timedelta
from openpyxl import load_workbook

PARAMETERS={
 'water_temp_c':'water_temperature_c','ph':'ph','do_mgL':'dissolved_oxygen_mg_l',
 'do_sat_pct':'dissolved_oxygen_saturation_pct','spc_uScm':'specific_conductivity_us_cm',
 'tds':'total_dissolved_solids_mg_l','mpn_final':'e_coli_mpn_100ml',
 'q_total_cfs_disp':'discharge_cfs','Coliform':'coliform','Thallium':'thallium',
 'Nitrate':'nitrate','Nitrite':'nitrite','Total Phosphorus':'total_phosphorus',
 'Calcium':'calcium','Magnesium':'magnesium','Sodium':'sodium','Potassium':'potassium',
 'Chloride':'chloride','Sulfate':'sulfate','Arsenic':'arsenic','Barium':'barium',
 'Cadmium':'cadmium','Copper':'copper','Iron':'iron','Lead':'lead','Manganese':'manganese',
 'Mercury':'mercury','Selenium':'selenium','Zinc':'zinc','Gross Alpha':'gross_alpha',
 'Gross Beta':'gross_beta','Uranium-238':'uranium_238','Radium-226':'radium_226'
}

def station_key(row):
 mlid=row.get('mlid')
 if mlid not in (None,'',0): return str(int(mlid))
 text=((row.get('site_name') or '')+' '+(row.get('notes') or '')).lower()
 if '4953550' in text or 'montezuma' in text: return '4953550'
 if 'castle' in text and ('trib' in text or 'side' in text): return '4958073'
 return None

def numeric(value):
 if value in (None,''): return None
 try: return float(value)
 except (TypeError,ValueError): return None

def main():
 ap=argparse.ArgumentParser();ap.add_argument('workbook');ap.add_argument('output');args=ap.parse_args()
 ws=load_workbook(args.workbook,read_only=True,data_only=True)['data']
 headers=[cell.value for cell in next(ws.iter_rows())];records=[];mst=timezone(timedelta(hours=-7))
 for source_row,row in enumerate(ws.iter_rows(min_row=2,values_only=True),2):
  item=dict(zip(headers,row));sampled=item.get('sample_dt_mst');mlid=station_key(item)
  if not mlid or not isinstance(sampled,datetime): continue
  measurements={}
  for source,target in PARAMETERS.items():
   value=numeric(item.get(source))
   if value is not None: measurements[target]=value
  if not measurements: continue
  records.append({'mlid':mlid,'sampled_at':sampled.replace(tzinfo=mst).isoformat(),
   'trip_id':item.get('trip_id'),'source_site_name':item.get('site_name'),
   'notes':item.get('notes'),'measurements':measurements,
   'source_document':'Southeast Utah Water Quality Monitoring V2(1).xlsx',
   'source_row_number':source_row})
 with open(args.output,'w',encoding='utf-8') as f: json.dump(records,f,separators=(',',':'),ensure_ascii=False)
 print(json.dumps({'records':len(records),'sites':len({r['mlid'] for r in records}),
  'results':sum(len(r['measurements']) for r in records)},indent=2))

if __name__=='__main__': main()
